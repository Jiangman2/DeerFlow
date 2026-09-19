#!/usr/bin/env python3
"""Maintenance helper: extract the 33 source rows from the reference workbook.

This is not part of report generation. Run it only when the finance department
replaces the authoritative workbook and the catalog must be refreshed.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("output")
    args = parser.parse_args()

    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required only for catalog maintenance") from exc

    workbook = openpyxl.load_workbook(args.workbook, read_only=False, data_only=False)
    sheet = workbook["示例汇总"]
    rows = []
    for row_number in range(2, 35):
        marker = sheet.cell(row_number, 1).value
        rows.append(
            {
                "source_row": row_number,
                "source_selection_mark": "selected"
                if marker == "R"
                else "not_selected",
                "title": sheet.cell(row_number, 2).value,
                "profile_label": sheet.cell(row_number, 3).value,
                "example": sheet.cell(row_number, 4).value,
            }
        )
    payload = {
        "version": "1.0",
        "source_file": Path(args.workbook).name,
        "selection_note": "The R/£ values are Wingdings 2 selection marks in the demonstration workbook, not mandatory/optional legal classifications.",
        "row_count": len(rows),
        "rows": rows,
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(f"CATALOG_EXTRACTED rows={len(rows)} output={output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
