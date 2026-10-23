"""Create the OA-facing V1.2 specification using bundled python-docx."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import argparse
import hashlib
import json
import zipfile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

parser = argparse.ArgumentParser()
parser.add_argument('--output', type=Path, required=True)
parser.add_argument('--reference', type=Path, default=Path('C:/Users/15285/Documents/deer-flow/deer-flow/docs/集团OA制度文生图数据服务接口规范说明文档_V1.1.docx'))
args = parser.parse_args()
reference = args.reference.read_bytes()
assert hashlib.sha256(reference).hexdigest() == '7103256916efb1988feba586840065ccc9b0b31b58565f55b2f159e634e774bf'
doc = Document(BytesIO(reference))
original = list(doc.element.body)
patterns = {'body': deepcopy(original[13]), 'h1': deepcopy(original[12]),
            'h2': deepcopy(original[15]), 'code': deepcopy(original[44])}
source_table = deepcopy(doc.tables[0]._tbl)
cell_patterns = [deepcopy(doc.tables[0].rows[i].cells[0]._tc) for i in range(3)]
for e in original[12:]:
    if e.tag != qn('w:sectPr'): doc.element.body.remove(e)
for e in original[:6]:
    for t in e.xpath('.//w:t'):
        t.text = t.text.replace('V1.1', 'V1.2').replace('2026-09-10', '2026-10-23')

def set_cell_text(cell, value):
    para=cell.paragraphs[0]
    rpr=deepcopy(para.runs[0]._r.rPr)
    para.clear(); para.add_run(value)._r.insert(0,rpr)

info=doc.tables[0]
set_cell_text(info.cell(2,1),'V1.2')
set_cell_text(info.cell(6,1),'提交、查询和 v1 下载已启用；v2 多文件及回调以平台启用成功通知为准。')
history=doc.tables[1]
row=deepcopy(history.rows[-1]._tr); history._tbl.append(row)
for cell,text in zip(history.rows[-1].cells,['V1.2','2026-10-23','新增显式 v2、多文件清单、逐文件下载与 v2 回调签名；保留 v1 兼容。']):
    set_cell_text(cell,text)

def styled(text, kind):
    e=deepcopy(patterns[kind]); v=Paragraph(e,doc)
    rpr=deepcopy(v.runs[0]._r.rPr)
    v.clear(); v.add_run(text)._r.insert(0,rpr)
    doc.element.body.insert(len(doc.element.body)-1,e)
    return v

def p(text,bold=False):
    v=styled(text,'body')
    if bold:v.runs[0].bold=True
    return v

def h(text):return styled(text,'h1')
def h2(text):return styled(text,'h2')
def page():pass  # Source uses continuous body flow; only its cover is a separate page.

def code(text,size=8):
    # Clone the original single-paragraph, gray, Consolas code-block component.
    v=styled(text.strip('\n'),'code')
    v.paragraph_format.keep_together=True
    return v

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t._tbl.remove(t._tbl.tblPr);t._tbl.insert(0,deepcopy(source_table.tblPr))
    scale=(11909-2*1037)/567.0/sum(widths)
    widths=[w*scale for w in widths]
    for j,w in enumerate(widths):t.columns[j].width=Cm(w)
    for idx,vals in enumerate([headers]+rows):
        cells=t.rows[0].cells if idx==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(cells,vals)):
            pattern=cell_patterns[0 if idx==0 else (1 if idx%2 else 2)]
            cell._tc.remove(cell._tc.tcPr);cell._tc.insert(0,deepcopy(pattern.tcPr))
            cell.width=Cm(widths[j])
            para=cell.paragraphs[0]
            if para._p.pPr is not None:para._p.remove(para._p.pPr)
            para._p.insert(0,deepcopy(pattern.xpath('./w:p/w:pPr')[0]))
            if idx and j:para.alignment=WD_ALIGN_PARAGRAPH.LEFT
            para.add_run(text)._r.insert(0,deepcopy(pattern.xpath('./w:p/w:r/w:rPr')[0]))
    v=doc.add_paragraph();v.paragraph_format.space_after=Pt(5);v.paragraph_format.line_spacing=1
    return t

h('1 适用范围与版本变化')
p('本文件在 2026-09-10 发布的 V1.1 文档基础上，完整说明 OA 端接入多文件结果所需的提交、查询、下载及回调协议。OA 改造后须显式选择 v2，并完整保存和展示结果清单中的全部独立 JSON 流程文件。')
p('生效说明：本次计划开启 OA 接口、多文件 v2 及主动回调。实际开放时间以平台启用成功通知为准；本文不表示 OA 业务联调已经通过。文档版本 V1.2 与请求中的 result_protocol=v2 是两个不同概念。',True)
h2('1.1 主要变化与兼容规则')
table(['事项','V1.1 原调用即协议 v1','本次协议 v2'],[
 ['提交地址与鉴权','固定 HTTPS 地址和专用机器密钥','保持不变'],
 ['版本选择','不传版本或传 v1','新增 result_protocol=v2'],
 ['结果交付','唯一 JSON 文件','1 至 10 份独立 JSON 文件的清单'],
 ['下载方式','/{job_id}/result','/{job_id}/results/{result_id}'],
 ['成功回调','manifest 加 result_file','manifest 加清单对应的各文件部分'],
 ['回调签名','时间戳 原始清单 单文件摘要','v2 前缀 时间戳 原始清单']
],[3.1,6.2,8.1])
p('不传版本的老调用继续按 v1 处理。v1 多文件输出仍不满足单文件契约；需要多份结果时必须改为 v2。平台不会把多份文件合并，也不会默认只返回第一份。回调接收端应同时兼容 v1 和 v2，以便存量调用继续使用。')
h2('1.2 接口入口')
code('BASE = https://cnoocclaw.cnooc.com.cn\nJOBS = /api/v1/integrations/oa/policy-flow/jobs',9)
p('后文的 JOBS 为上述路径前缀，示例中的键、任务 ID、摘要和文件名均为占位或虚构数据。OA 不需要直接调用 Gateway、模型服务或浏览器认证接口。')

page(); h('2 通用协议与任务提交')
p('使用 HTTPS 和 UTF-8。文件使用 multipart/form-data 原始文件部分，不使用 Base64；Content-Type 的 boundary 由客户端自动生成。响应时间是带时区的字符串，签名时间戳使用 Unix 秒。')
table(['请求头','要求'],[
 ['Authorization','Bearer <INGEST_KEY>。提交、查询和下载均使用同一 OA 专用接口密钥。'],
 ['Idempotency-Key','仅提交必填。8–200 个可见 ASCII 字符，不含空格；建议 policy:<request_id>。'],
 ['X-OA-Integration-Key','兼容的替代鉴权头；建议统一使用 Bearer，不同时传入两套不同密钥。']
],[4.4,13.0])
h2('2.1 提交地址和表单')
code('POST https://cnoocclaw.cnooc.com.cn/api/v1/integrations/oa/policy-flow/jobs',8.1)
table(['字段','类型与必填','约束'],[
 ['document','file 必填','一份非空 DOC、DOCX 或 PDF；不超过 10 MiB，并受平台上传限制及安全检查约束。'],
 ['request_id','string 必填','OA 业务请求号，1–128 字符，去除首尾空白后不得为空，不得含控制字符。'],
 ['key_fields','string 必填','严格 JSON 对象的 UTF-8 文本，最大 64 KiB；拒绝重复键、NaN、Infinity 等。'],
 ['result_protocol','string 可选','v2 为多文件；不传或 v1 为旧单文件；其他值返回 422。']
],[3.6,3.4,10.4])
h2('2.2 业务字段沿用原文档')
table(['key_fields 内字段','要求','含义'],[
 ['documentCode','必填 不超过128字符','制度编码，例如 CG-02-03'],
 ['documentName','必填 不超过256字符','制度完整名称'],
 ['version','必填 不超过64字符','制度版本，例如 V1.0'],
 ['organizationCode','选填 不超过64字符','所属组织编码']
],[4.3,5.4,7.7])
p('以上是沿用 V1.1 的 OA 业务约定；OA 应在提交前校验。服务端协议层校验 JSON 对象结构和大小，回调原样返回该业务对象，不把 key_fields 当作生成指令，也不允许其中指定回调地址、用户或智能体。')
h2('2.3 幂等处理')
p('同一幂等键、request_id、key_fields 对象、文件内容和协议版本重复提交，返回原 job_id，不重复启动任务。相同键对应不同内容或协议时返回 409。新增业务请求或由 v1 改为 v2 时使用新的幂等键。网络超时不代表提交失败，先复用原键核对，不能换键盲目重发。')

page();h('3 提交示例与受理响应')
p('以下为 Bash/curl 示例。INGEST_KEY 从安全渠道取得并在运行环境中提供，示例不含真实密钥。删除旧文档示例行首的“+”符号，不要把它作为命令字符。')
code(r'''curl --fail-with-body --connect-timeout 10 --max-time 60 \
  'https://cnoocclaw.cnooc.com.cn/api/v1/integrations/oa/policy-flow/jobs' \
  -H "Authorization: Bearer ${INGEST_KEY}" \
  -H 'Idempotency-Key: policy:OA-TEST-20261023-001:v2' \
  --form-string 'request_id=OA-TEST-20261023-001' \
  --form-string 'result_protocol=v2' \
  --form-string 'key_fields={"documentCode":"CG-02-03","documentName":"采购管理制度","version":"V1.0","organizationCode":"10001"}' \
  -F 'document=@制度文件.docx;type=application/vnd.openxmlformats-officedocument.wordprocessingml.document'
''',8)
p('首次受理返回 HTTP 202。202 仅表示任务已持久化，不能作为生成成功或 OA 已收文件的依据。重复提交也返回原任务当前状态，可能已经为 succeeded 或 failed。')
code(json.dumps({
 'job_id':'54a4dd5b-bc52-42f6-a5ed-559e7162fc15','request_id':'OA-TEST-20261023-001',
 'status':'queued','callback_status':'not_ready','created_at':'2026-10-23 02:00:00+00:00',
 'updated_at':'2026-10-23 02:00:00+00:00','completed_at':None,'delivered_at':None,
 'error_code':None,'error_detail':None,'callback_error':None,'result_filename':None,
 'result_sha256':None,'result_protocol':'v2','results':[]},ensure_ascii=False,indent=2),8.6)
h2('OA 收到受理响应后的处理')
p('保存 request_id、job_id、幂等键和版本。后续以 job_id 查询或接收回调。不能用“文件数量已非零”替代 status=succeeded；任务失败时应记录 error_code 和 error_detail，不把目录中尚存的中间文件作为交付。')
p('建议每 10–30 秒查询一次；单次提交建议超时 60 秒。当前业务运行上限为 1800 秒，排队与重试可能额外耗时。建议设定可配置的等待窗口，并在超时后保留任务号继续查询或人工排查，不据客户端等待超时直接重建任务。')

page();h('4 状态查询与响应字段')
code('GET {BASE}{JOBS}/{job_id}\nAuthorization: Bearer <INGEST_KEY>',9)
p('提交和状态查询使用同一响应结构。v1 响应保持旧字段；v2 增加 result_protocol 和 results。状态查询不回传 key_fields，OA 应保留提交时的业务关联信息；回调 manifest 会携带 key_fields。')
table(['字段','类型','含义'],[
 ['job_id / request_id','string','平台任务 ID / OA 业务请求号'],
 ['status','string','queued / launching / running / succeeded / failed'],
 ['callback_status','string','not_ready / pending / delivered / exhausted / disabled'],
 ['created_at / updated_at','string','创建时间 / 最近状态更新时间，含时区'],
 ['completed_at / delivered_at','string 或 null','生成完成时间 / OA 回调交付时间'],
 ['error_code / error_detail','string 或 null','生成失败代码 / 安全处理后的说明'],
 ['callback_error','string 或 null','最近一次回调失败说明'],
 ['result_filename / result_sha256','string 或 null','v1 单文件结果；v2 为 null'],
 ['result_protocol','string','v2 任务返回 v2；v1 响应可无此字段'],
 ['results','array','v2 文件清单；受理时为空，成功后含全部结果']
],[6,3.2,8.2])
table(['状态','OA 应如何处理'],[
 ['queued / launching / running','已受理等待处理 / 正在准备与启动 / 智能体运行中。继续查询或等待回调。'],
 ['succeeded','结果已校验并保存。处理完整清单；回调失败不改变此状态。'],
 ['failed','生成失败。查看错误信息，不把它当作有效结果任务。'],
 ['not_ready / pending','尚未进入回调 / 等待回调或等待重试。'],
 ['delivered','OA 接口已返回 2xx。OA 仍需确保自身入库和渲染正确。'],
 ['exhausted','回调出现不可重试错误或重试耗尽；可查询并下载补偿。'],
 ['disabled','生成完成时回调关闭；OA 应查询并下载。历史 disabled 不会因开启回调自动补发。']
],[5.4,12])
p('两组状态含义独立。例如 status=succeeded、callback_status=exhausted 表示生成成功但推送未完成，不应重新生成整份制度。')

page();h('5 多文件清单与逐文件下载')
table(['results 中字段','含义'],[
 ['id','结果 ID，例如 result_001。用于下载，禁止由文件名或序号自行猜测。'],
 ['part_name','回调 multipart 中对应的文件字段名。'],
 ['filename','展示文件名，可含中文；OA 存储时需自行做安全文件名处理。'],
 ['bytes','该文件原始字节数，不是字符数。'],
 ['sha256','原始文件字节的 SHA-256，小写 64 位十六进制。']
],[4.1,13.3])
p('成功响应的关键字段示例如下。摘要字符串是说明性占位，接入时必须使用接口返回的实际值。')
code('''{
  "job_id": "54a4dd5b-bc52-42f6-a5ed-559e7162fc15",
  "status": "succeeded", "callback_status": "pending",
  "result_protocol": "v2",
  "result_filename": null, "result_sha256": null,
  "results": [
    {"id":"result_001", "part_name":"result_001",
     "filename":"采购申请流程.json", "bytes":1234,
     "sha256":"<文件一的64位小写十六进制摘要>"},
    {"id":"result_002", "part_name":"result_002",
     "filename":"采购审批流程.json", "bytes":2345,
     "sha256":"<文件二的64位小写十六进制摘要>"}
  ]
}''',8.5)
code('''GET {BASE}{JOBS}/{job_id}/results/{result_id}
Authorization: Bearer <INGEST_KEY>

HTTP/1.1 200 OK
Content-Type: application/json
Content-Disposition: attachment; filename*=UTF-8''<URL编码文件名>
Cache-Control: private, no-store
X-Content-Type-Options: nosniff
ETag: "<该文件sha256>"''',8.6)
p('OA 必须遍历 results 下载全部文件，以 job_id + id 关联、幂等保存；按原始字节核对 bytes 和 sha256 后再解析 UTF-8 JSON。不要先格式化或重新序列化 JSON 再计算摘要。每份文件独立交给 JointJS 加载，不能只展示第一份或拼接多个 JSON 文本。')
p('当前数量上限 10 份、单份 10 MiB、整批 50 MiB。平台对整批校验，不接受非法 JSON、重复结果、重名、越界路径或超限；一份无效会导致整批验收失败。JSON 格式通过仍需 OA 判断流程业务正确性。')
p('v2 使用旧 /{job_id}/result 返回 409。v1 继续使用该旧地址；v1 任务调用多文件下载也返回 409。结果缺失或快照损坏返回 410，不能自动视作成功。')

page();h('6 主动回调接口与清单')
p('平台在任务成功或失败后 POST 到双方已配置的固定 HTTPS 回调地址。现有目标主机为 grc-web.apps.shdev.ocpx.cnooc，具体路径保留服务器现有配置；如需变更，应由双方管理员通过配置流程更新，不能在任务提交中动态指定。平台不跟随重定向。')
table(['请求头','说明'],[
 ['Content-Type','multipart/form-data；每个文件为独立部分。'],
 ['Idempotency-Key','oa-policy-flow-callback:<job_id>，用于 OA 端去重。'],
 ['X-OA-Callback-Timestamp','Unix 秒级时间戳，字符串原值参与验签。'],
 ['X-OA-Callback-Signature','sha256=<HMAC-SHA256 小写十六进制签名>。'],
 ['Authorization','仅在双方额外配置回调 Bearer 时携带；当前使用 HMAC，无回调 Bearer。']
],[6.2,11.2])
p('成功回调包含一个名为 manifest 的 JSON 文本部分，以及 results 清单中每项 part_name 对应的一个文件部分。manifest 为文本表单部分，不附文件名，Content-Type 为 application/json。文件部分 Content-Type 同为 application/json。v2 不再使用 result_file 字段。')
code('''{
  "job_id": "54a4dd5b-bc52-42f6-a5ed-559e7162fc15",
  "request_id": "OA-TEST-20261023-001",
  "status": "succeeded",
  "key_fields": {"documentCode":"CG-02-03",
    "documentName":"采购管理制度", "version":"V1.0",
    "organizationCode":"10001"},
  "completed_at": "2026-10-23 02:20:00+00:00",
  "error_code": null, "error_detail": null,
  "result_protocol": "v2",
  "results": [
    {"id":"result_001", "part_name":"result_001",
     "filename":"采购申请流程.json", "bytes":1234,
     "sha256":"<文件一的64位小写十六进制摘要>"},
    {"id":"result_002", "part_name":"result_002",
     "filename":"采购审批流程.json", "bytes":2345,
     "sha256":"<文件二的64位小写十六进制摘要>"}
  ]
}''',8.1)
p('以上 manifest 对应两个文件部分 result_001、result_002。OA 必须检查文件部分与清单一一对应，拒绝缺失、重复或额外文件，并逐份比对文件名、字节数和摘要。')
p('失败回调保留相同任务和业务字段，status=failed，error_code/error_detail 给出失败说明，result_protocol=v2、results=[]，不附带任何结果文件。v2 清单不含单值 result_filename/result_sha256。')

page();h('7 HMAC 验签与回调幂等')
p('沿用原 V1.1 的 HMAC 必选约定和既有共享密钥；接口访问密钥与回调签名密钥是不同凭据。时间窗口建议前后 300 秒，由双方运维确认并保持时钟同步。验证原始 manifest UTF-8 文本，不能重新排序、去空格或转义后再验签。')
table(['协议','参与 HMAC 的完整消息'],[
 ['v2','UTF8("v2." + timestamp + "." + manifest原始文本)'],
 ['v1','UTF8(timestamp + "." + manifest原始文本 + "." + 单文件sha256)；失败时最后的sha256为空字符串']
],[2,15.4])
p('签名头为 sha256= 加摘要的小写十六进制。OA 可读取原始清单中的 result_protocol 选择算法，但只有通过签名验证后才能信任其中业务数据。v1 无 result_protocol 字段；收到未知协议应拒绝。')
p('以下 Python 示例仅示范验签。manifest_text 必须直接来自收到的文本部分；正式接收端还需执行严格 JSON、文件校验及幂等持久化。')
code('''import hashlib, hmac, json, time

def verify_signature(manifest_text, timestamp, signature, secret):
    try:
        ts = int(timestamp)
        if abs(time.time() - ts) > 300:
            return False
        obj = json.loads(manifest_text)
        protocol = obj.get("result_protocol", "v1")
        if protocol == "v2":
            message = "v2." + timestamp + "." + manifest_text
        elif protocol == "v1":
            digest = obj.get("result_sha256") or ""
            message = timestamp + "." + manifest_text + "." + digest
        else:
            return False
        expected = "sha256=" + hmac.new(
            secret.encode("utf-8"), message.encode("utf-8"),
            hashlib.sha256).hexdigest()
        return hmac.compare_digest(expected, signature)
    except (ValueError, TypeError, AttributeError, OverflowError):
        return False''',8.4)
h2('接收顺序与返回规则')
p('先限制请求大小并解析 multipart，再校验时间和签名、严格 JSON 及文件对应关系、原始字节摘要，最后按 job_id + result_id 原子或可恢复地持久化。全部文件成功保存后返回任意 2xx；平台不读取响应正文，不要求 code/msg/data 包装。')
p('对已完整保存的重复回调直接返回 2xx，不重复建立业务流程。部分接收失败不得提前返回成功；重试应补齐同一任务。数据库状态与文件落盘应协调，避免返回 2xx 后实际只保存部分结果。失败任务通知按 job_id 幂等保存错误信息。')
p('网络错误及 408、409、425、429、5xx 有界重试，当前最多 8 次；其他状态不会按此规则继续重试。不要用 409 表示“此任务早已完整接收”，该场景应返回 2xx。重试会刷新时间戳和签名，业务幂等键仍保持不变。')

page();h('8 错误处理与 OA 改造验收')
p('接口直接使用 HTTP 状态码，不套 code/msg/data 统一外层。路由错误通常为 {"detail":"说明"}；框架参数校验的 detail 也可能是数组，OA 应兼容并保留脱敏诊断信息。')
table(['HTTP 状态','场景与处理'],[
 ['200 / 202','查询或下载成功 / 提交受理。202 不是生成成功。'],
 ['401','机器接口密钥无效；核对环境及密钥，不循环重试。'],
 ['404','任务或结果 ID 不存在；核对已保存的 job_id 和清单 ID。'],
 ['409','幂等冲突、任务未成功或下载协议不匹配；根据 detail 分别处理。'],
 ['410','已接受的结果无法读取；联系平台排查，不能伪造或静默略过文件。'],
 ['413 / 415','上传或字段超限 / 文件类型不支持。修正请求后按新业务内容提交。'],
 ['422','表单、版本、幂等键或 JSON 格式无效；修正参数。'],
 ['500','持久化等内部异常；保留原幂等键核对，记录时间与请求号。'],
 ['503','OA 服务不可用，或 v2 开关未启用；核对平台开放通知。此状态不等同于模型服务报错。']
],[2.7,14.7])
h2('OA 端必须完成的改造')
p('1. 提交新增 result_protocol=v2，保留原鉴权、幂等和业务字段；不改存量 v1 任务的版本。\n2. 任务数据结构增加结果列表，按 job_id + result_id 保存多个文件及各自摘要。\n3. 状态查询成功后遍历全部结果，用新下载路径获取每份文件。\n4. 回调端识别 v1/v2，分别使用对应 HMAC 算法和文件字段，并逐份校验、幂等保存。\n5. 前端按清单展示全部独立流程，逐个验证 JointJS 加载和制度业务含义。')
h2('开放后建议由 OA 执行的验收')
p('至少验证一个两份或三份结果任务、一个原单文件 v1 任务、同键同内容重复提交、同键不同内容冲突、错误密钥、回调重复接收、部分接收失败后的重试，以及回调失败后的逐文件下载补偿。回调数据返回 2xx 后应核对 OA 实际保存数量和摘要，不能只依据 HTTP 成功判定业务验收。')
p('本次不自动重放历史失败任务，也不把历史错误改成成功。历史 callback_status=disabled 的任务不会因开启回调自动补发。调用中出现异常时，请提供北京时间、request_id、job_id、协议版本、HTTP 状态、error_code、callback_status 及脱敏日志，不提供接口密钥、HMAC 密钥或制度正文。')

args.output.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.title='集团 OA 制度文生图数据服务接口规范说明文档 V1.2'
doc.core_properties.subject='兼容 v1 的 v2 多文件接口改造说明'
doc.core_properties.author='智能体平台'
buffer=BytesIO();doc.save(buffer)
with zipfile.ZipFile(BytesIO(reference)) as src, zipfile.ZipFile(buffer) as edited:
    changes={name:edited.read(name) for name in ['word/document.xml','docProps/core.xml']}
    changes['word/footer1.xml']=src.read('word/footer1.xml').replace(b'V1.1',b'V1.2')
    settings=etree.fromstring(src.read('word/settings.xml'))
    field=settings.find(qn('w:updateFields'))
    if field is None:field=etree.SubElement(settings,qn('w:updateFields'))
    field.set(qn('w:val'),'true')
    changes['word/settings.xml']=etree.tostring(settings,encoding='UTF-8',xml_declaration=True,standalone=True)
    with zipfile.ZipFile(args.output,'w',compression=zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():out.writestr(item,changes.get(item.filename,src.read(item.filename)))
assert args.reference.read_bytes()==reference
print(json.dumps({'output':str(args.output),'bytes':args.output.stat().st_size,'reference_sha256':hashlib.sha256(reference).hexdigest(),'preserved_package_parts':14},ensure_ascii=True))
