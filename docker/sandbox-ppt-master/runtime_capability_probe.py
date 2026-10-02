#!/usr/bin/env python3
"""Production-readiness probes for the isolated ppt-master Agent image.

The probe deliberately separates deterministic local tooling checks from
operator-owned network capabilities.  It never prints secret values and writes
all canary artifacts below a temporary directory that is removed on exit.
"""

from __future__ import annotations

import argparse
import importlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib import error, request

try:
    import tomllib
except ModuleNotFoundError:  # Python 3.10 production Sandbox compatibility.
    import tomli as tomllib


SKILL_ROOT = Path(
    os.environ.get("PPT_MASTER_SOURCE_ROOT", "/opt/ppt-master-release-skill")
).resolve()
EXPECTED_VERSION = "6.6.0"
EXPECTED_MODEL = "DeepSeek-V4-Flash"
EXPECTED_PROVIDER = "cnooc_deepseek"
EXPECTED_RELAY_URL = "http://gateway:8001/api/codex/deepseek/v1"
EXPECTED_IMAGE_MODEL = "qwen-image-2.0-pro"
EXPECTED_TTS_MODEL = "qwen3-tts-flash"
DEFAULT_TTS_BASE_URL = "https://dashscope.aliyuncs.com"
AGENT_MARKER = "PPT_MASTER_AGENT_LLM_CANARY_OK"


def fail(detail: str) -> None:
    raise RuntimeError(detail)


def require(condition: object, detail: str) -> None:
    if not condition:
        fail(detail)


def run(command: list[str], *, timeout: int = 180, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        command,
        cwd=cwd,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=timeout,
        check=False,
        env=dict(os.environ),
    )
    if result.returncode != 0:
        output = (result.stdout or "").strip()
        if len(output) > 4000:
            output = output[-4000:]
        fail("command_failed=%r exit_code=%d output=%s" % (command, result.returncode, output))
    return result


def import_required_modules() -> None:
    modules = (
        "yaml",
        "pptx",
        "xlsxwriter",
        "pathops",
        "uharfbuzz",
        "edge_tts",
        "fitz",
        "mammoth",
        "markdownify",
        "ebooklib",
        "nbconvert",
        "openpyxl",
        "PIL",
        "numpy",
        "requests",
        "bs4",
        "curl_cffi",
        "google.genai",
        "flask",
        "cairosvg",
        "duckdb",
        "playwright.sync_api",
        "docx",
        "pdfplumber",
        "pypdf",
        "reportlab",
        "markdown",
        "chardet",
        "statsmodels",
        "sklearn",
        "tomli" if sys.version_info < (3, 11) else "tomllib",
    )
    missing: list[str] = []
    for name in modules:
        try:
            importlib.import_module(name)
        except Exception as exc:  # pragma: no cover - executed in release image
            missing.append("%s:%s" % (name, type(exc).__name__))
    require(not missing, "missing_python_modules=" + ",".join(missing))


def skill_version() -> str:
    body = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8")
    match = re.search(r'^\s{2}version:\s*["\']([^"\']+)["\']\s*$', body, re.MULTILINE)
    return match.group(1) if match else ""


def local_toolchain_probe() -> None:
    require((SKILL_ROOT / "SKILL.md").is_file(), "ppt_master_skill_missing")
    require(skill_version() == EXPECTED_VERSION, "ppt_master_skill_version_mismatch")
    for relative in (
        "scripts/image_gen.py",
        "scripts/image_search.py",
        "scripts/notes_to_audio.py",
        "scripts/pptx_delivery_check.py",
        "workflows/generate-pptx.md",
        "workflows/edit-native-pptx.md",
    ):
        require((SKILL_ROOT / relative).is_file(), "ppt_master_path_missing=" + relative)
    import_required_modules()

    for command in (
        "libreoffice",
        "pandoc",
        "ffmpeg",
        "ffprobe",
        "fc-match",
        "rsvg-convert",
        "websocat",
        "codex",
        "deerflow-markdown-to-pdf",
    ):
        require(shutil.which(command), "required_command_missing=" + command)

    with tempfile.TemporaryDirectory(prefix="ppt-master-local-canary-") as raw_temp:
        root = Path(raw_temp)
        from pptx import Presentation
        import cairosvg
        from playwright.sync_api import sync_playwright

        deck = root / "canary.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "PPT Master 生产工具链验收"
        presentation.save(deck)
        require(deck.stat().st_size > 0, "python_pptx_canary_empty")

        office_out = root / "office"
        office_out.mkdir()
        run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(office_out),
                str(deck),
            ],
            timeout=180,
        )
        require((office_out / "canary.pdf").stat().st_size > 0, "libreoffice_pdf_canary_empty")

        markdown = root / "canary.md"
        html = root / "canary.html"
        markdown.write_text("# PPT Master\n\n生产工具链验收。\n", encoding="utf-8")
        run(["pandoc", str(markdown), "-o", str(html)])
        require("PPT Master" in html.read_text(encoding="utf-8"), "pandoc_canary_invalid")

        audio = root / "canary.wav"
        run(
            [
                "ffmpeg",
                "-nostdin",
                "-loglevel",
                "error",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=880:duration=0.25",
                "-y",
                str(audio),
            ]
        )
        require(audio.stat().st_size > 0, "ffmpeg_audio_canary_empty")

        png = root / "canary.png"
        cairosvg.svg2png(
            bytestring=(
                '<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64">'
                '<rect width="64" height="64" fill="#005bbb"/></svg>'
            ).encode("utf-8"),
            write_to=str(png),
        )
        require(png.stat().st_size > 0, "cairosvg_canary_empty")

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True, args=["--no-sandbox"])
            page = browser.new_page()
            page.set_content("<title>ppt-master-ready</title><p>ready</p>")
            require(page.title() == "ppt-master-ready", "chromium_canary_invalid")
            browser.close()

        font = run(["fc-match", "Noto Sans CJK SC"], timeout=30).stdout
        require(font.strip(), "cjk_font_canary_empty")

        document_markdown = root / "document.md"
        document_pdf = root / "document.pdf"
        document_markdown.write_text("# 中文 PDF 验收\n\n沙箱文档运行时正常。\n", encoding="utf-8")
        run(["deerflow-markdown-to-pdf", str(document_markdown), str(document_pdf)], timeout=180)
        import fitz

        with fitz.open(document_pdf) as pdf:
            extracted = "\n".join(page.get_text() for page in pdf)
        require("中文 PDF 验收" in extracted, "cjk_pdf_text_roundtrip_failed")

        from docx import Document

        document_docx = root / "document.docx"
        word = Document()
        word.add_paragraph("python-docx 沙箱验收")
        word.save(document_docx)
        require(document_docx.stat().st_size > 0, "python_docx_canary_empty")

    print("ppt_master_local_toolchain=PASS")
    print("ppt_master_office_pdf_conversion=PASS")
    print("ppt_master_chromium_rendering=PASS")
    print("ppt_master_audio_video_toolchain=PASS")
    print("sandbox_document_runtime=PASS python_docx=yes cjk_pdf=yes")


def required_env(name: str) -> str:
    value = os.environ.get(name, "").strip()
    require(value, "required_environment_missing=" + name)
    return value


def configuration_probe() -> None:
    relay_key = required_env("PPT_MASTER_RESPONSES_RELAY_KEY")
    dashscope_key = required_env("DASHSCOPE_API_KEY")
    sandbox_key = required_env("SANDBOX_API_KEY")
    require(len(relay_key) >= 16, "relay_key_too_short")
    require(len(dashscope_key) >= 8, "dashscope_key_too_short")
    require(len(sandbox_key) >= 16, "sandbox_key_too_short")

    config = tomllib.loads(required_env("CODEX_CONFIG_TOML"))
    require(config.get("model") == EXPECTED_MODEL, "codex_model_mismatch")
    require(config.get("model_provider") == EXPECTED_PROVIDER, "codex_provider_mismatch")
    providers = config.get("model_providers")
    provider = providers.get(EXPECTED_PROVIDER) if isinstance(providers, dict) else None
    require(isinstance(provider, dict), "codex_provider_section_missing")
    require(str(provider.get("base_url") or "").rstrip("/") == EXPECTED_RELAY_URL, "codex_relay_url_mismatch")
    require(provider.get("env_key") == "PPT_MASTER_RESPONSES_RELAY_KEY", "codex_relay_env_key_mismatch")
    require(provider.get("wire_api") == "responses", "codex_wire_api_mismatch")
    require(os.environ.get("PPT_MASTER_ENTERPRISE_RELAY_REQUIRED", "").strip().lower() == "true", "enterprise_relay_not_required")

    require(required_env("IMAGE_BACKEND") == "qwen", "image_backend_mismatch")
    require(required_env("QWEN_MODEL") == EXPECTED_IMAGE_MODEL, "image_model_mismatch")
    require(required_env("QWEN_BASE_URL").startswith("https://"), "image_base_url_not_https")
    require(required_env("DASHSCOPE_TTS_MODEL") == EXPECTED_TTS_MODEL, "tts_model_mismatch")
    require(required_env("DASHSCOPE_TTS_VOICE_MALE"), "tts_male_voice_missing")
    require(required_env("DASHSCOPE_TTS_VOICE_FEMALE"), "tts_female_voice_missing")
    tts_base = os.environ.get("QWEN_TTS_BASE_URL", "").strip()
    require(tts_base.rstrip("/") == DEFAULT_TTS_BASE_URL, "qwen_tts_base_url_mismatch")

    keyed_sources = sum(bool(os.environ.get(name, "").strip()) for name in ("PEXELS_API_KEY", "PIXABAY_API_KEY"))
    print(
        "ppt_master_model_configuration=PASS "
        "llm=%s image=%s tts=%s" % (EXPECTED_MODEL, EXPECTED_IMAGE_MODEL, EXPECTED_TTS_MODEL)
    )
    print("ppt_master_web_image_sources=PASS keyed_sources=%d zero_key_fallbacks=2" % keyed_sources)


def opener() -> request.OpenerDirector:
    return request.build_opener(request.ProxyHandler({}))


def json_request(url: str, *, method: str = "GET", payload: dict | None = None, token: str = "", timeout: int = 30) -> dict:
    body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    headers = {"Accept": "application/json"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = "Bearer " + token
    call = request.Request(url, data=body, headers=headers, method=method)
    try:
        with opener().open(call, timeout=timeout) as response:
            value = json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:1000]
        fail("http_error=%d url=%s detail=%s" % (exc.code, url, detail))
    except error.URLError as exc:
        fail("http_transport_error url=%s reason=%s" % (url, exc.reason))
    require(isinstance(value, dict), "http_response_not_object=" + url)
    return value


def agent_live_probe() -> None:
    base = "http://127.0.0.1:%s" % os.environ.get("PPT_MASTER_AGENT_HOST_PORT", "8788")
    deadline = time.monotonic() + 120
    health: dict | None = None
    while time.monotonic() < deadline:
        try:
            health = json_request(base + "/health", timeout=3)
            if health.get("status") == "ok":
                break
        except Exception:
            health = None
        time.sleep(2)
    require(health and health.get("status") == "ok", "agent_host_health_timeout")
    provider = health.get("provider")
    require(isinstance(provider, dict), "agent_host_provider_metadata_missing")
    require(provider.get("model") == EXPECTED_MODEL, "agent_host_model_mismatch")
    require(provider.get("provider") == EXPECTED_PROVIDER, "agent_host_provider_mismatch")

    key = required_env("SANDBOX_API_KEY")
    task_id = "ppt-capability-%d" % int(time.time())
    task = json_request(
        base + "/v1/tasks",
        method="POST",
        token=key,
        payload={
            "task_id": task_id,
            "thread_id": "ppt-capability-canary",
            "resume": False,
            "prompt": (
                "This is a production readiness canary. Do not create files, do not run tools, "
                "and do not activate a Skill. Reply with exactly " + AGENT_MARKER + " and nothing else."
            ),
        },
        timeout=30,
    )
    require(task.get("task_id") == task_id, "agent_host_task_id_mismatch")
    deadline = time.monotonic() + 360
    while time.monotonic() < deadline:
        task = json_request(base + "/v1/tasks/" + task_id, token=key, timeout=15)
        status = task.get("status")
        if status in {"completed", "failed", "cancelled"}:
            break
        time.sleep(2)
    require(task.get("status") == "completed", "agent_host_task_failed=" + str(task.get("error") or task.get("status")))
    final_message = str(task.get("final_message") or "")
    require(AGENT_MARKER in final_message, "agent_host_llm_marker_missing")
    usage = task.get("usage")
    require(isinstance(usage, dict), "agent_host_token_usage_missing")
    require(int(usage.get("input_tokens") or 0) > 0, "agent_host_input_tokens_missing")
    require(int(usage.get("output_tokens") or 0) > 0, "agent_host_output_tokens_missing")
    require(int(usage.get("total_tokens") or 0) > 0, "agent_host_total_tokens_missing")
    require(int(usage.get("llm_call_count") or 0) >= 1, "agent_host_llm_call_count_missing")
    records = task.get("usage_records")
    require(isinstance(records, list) and records, "agent_host_usage_records_missing")
    print("ppt_master_agent_host_codex_live=PASS model=" + EXPECTED_MODEL)
    print("ppt_master_llm_relay_live=PASS wire_api=responses")
    print(
        "ppt_master_agent_host_token_usage=PASS input=%d output=%d total=%d calls=%d"
        % (
            int(usage["input_tokens"]),
            int(usage["output_tokens"]),
            int(usage["total_tokens"]),
            int(usage["llm_call_count"]),
        )
    )


def validate_image(path: Path) -> None:
    from PIL import Image

    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        require(image.width >= 64 and image.height >= 64, "generated_image_too_small")


def generated_image_file(root: Path, *, ignored: set[str]) -> Path:
    candidates = [
        path
        for path in root.iterdir()
        if path.is_file()
        and path.name not in ignored
        and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
    ]
    require(len(candidates) == 1, "unexpected_image_output_count=%d" % len(candidates))
    return candidates[0]


def image_live_probe() -> None:
    with tempfile.TemporaryDirectory(prefix="ppt-master-image-canary-") as raw_temp:
        root = Path(raw_temp)
        run(
            [
                sys.executable,
                str(SKILL_ROOT / "scripts" / "image_gen.py"),
                "A simple blue circular readiness badge on a white background, no text",
                "--backend",
                "qwen",
                "--model",
                EXPECTED_IMAGE_MODEL,
                "--aspect_ratio",
                "1:1",
                "--image_size",
                "512px",
                "--output",
                str(root),
                "--filename",
                "ppt_master_image_canary",
            ],
            timeout=600,
        )
        image = generated_image_file(root, ignored={"image_prompts.json", "image_prompts.md"})
        validate_image(image)
    print("ppt_master_image_generation_live=PASS model=" + EXPECTED_IMAGE_MODEL)


def tts_live_probe() -> None:
    scripts = str(SKILL_ROOT / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    from tts_backends import backend_qwen

    with tempfile.TemporaryDirectory(prefix="ppt-master-tts-canary-") as raw_temp:
        output = Path(raw_temp) / "canary.wav"
        backend_qwen.generate(
            "演示文稿旁白能力验收通过。",
            output,
            api_key=required_env("DASHSCOPE_API_KEY"),
            voice_id=required_env("DASHSCOPE_TTS_VOICE_FEMALE"),
            model=required_env("DASHSCOPE_TTS_MODEL"),
            language_type="Chinese",
            instructions=None,
            optimize_instructions=None,
            base_url=required_env("QWEN_TTS_BASE_URL"),
        )
        require(output.stat().st_size > 256, "tts_output_too_small")
        run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1", str(output)])
    print(
        "ppt_master_narration_tts_live=PASS model=%s voice=%s"
        % (required_env("DASHSCOPE_TTS_MODEL"), required_env("DASHSCOPE_TTS_VOICE_FEMALE"))
    )


def search_live_probe() -> None:
    if os.environ.get("PEXELS_API_KEY", "").strip():
        provider = "pexels"
    elif os.environ.get("PIXABAY_API_KEY", "").strip():
        provider = "pixabay"
    else:
        provider = "openverse"
    with tempfile.TemporaryDirectory(prefix="ppt-master-search-canary-") as raw_temp:
        root = Path(raw_temp)
        run(
            [
                sys.executable,
                str(SKILL_ROOT / "scripts" / "image_search.py"),
                "offshore wind turbine",
                "--provider",
                provider,
                "--filename",
                "ppt_master_search_canary.jpg",
                "--output",
                str(root),
                "--min-width",
                "64",
                "--min-height",
                "64",
            ],
            timeout=300,
        )
        output = root / "ppt_master_search_canary.jpg"
        require(output.is_file(), "image_search_output_missing")
        validate_image(output)
        manifest = root / "image_sources.json"
        require(manifest.is_file(), "image_search_manifest_missing")
        parsed = json.loads(manifest.read_text(encoding="utf-8"))
        require(parsed, "image_search_manifest_empty")
    print("ppt_master_web_image_search_live=PASS provider=" + provider)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--local", action="store_true")
    parser.add_argument("--config", action="store_true")
    parser.add_argument("--agent-live", action="store_true")
    parser.add_argument("--image-live", action="store_true")
    parser.add_argument("--tts-live", action="store_true")
    parser.add_argument("--search-live", action="store_true")
    parser.add_argument("--all", action="store_true")
    args = parser.parse_args()
    selected = args.all or any(
        (args.local, args.config, args.agent_live, args.image_live, args.tts_live, args.search_live)
    )
    if not selected:
        parser.error("select at least one probe mode")
    try:
        if args.all or args.local:
            local_toolchain_probe()
        if args.all or args.config:
            configuration_probe()
        if args.all or args.agent_live:
            agent_live_probe()
        if args.all or args.image_live:
            image_live_probe()
        if args.all or args.tts_live:
            tts_live_probe()
        if args.all or args.search_live:
            search_live_probe()
    except Exception as exc:
        print(
            "ppt_master_runtime_capabilities=FAIL detail=%s:%s"
            % (type(exc).__name__, str(exc)),
            file=sys.stderr,
        )
        return 1
    print("ppt_master_runtime_capabilities=PASS")
    print("ppt_master_native_powerpoint_video_export=WINDOWS_ONLY_NOT_REQUIRED_ON_LINUX")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
