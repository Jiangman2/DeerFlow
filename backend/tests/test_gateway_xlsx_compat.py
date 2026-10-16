"""Regression tests for lazy XLSX compatibility in Gateway conversion."""

from __future__ import annotations

import io
import os
import subprocess
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.dimensions import SheetFormatProperties

from deerflow.utils import file_conversion

REPO_ROOT = Path(__file__).resolve().parents[1]

_MARKER = "__deerflow_default_col_width_pt_compat__"


def _workbook_bytes() -> bytes:
    output = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "生产数据"
    worksheet["A1"] = "井号"
    worksheet["B1"] = "NB27-1-1"
    workbook.save(output)
    return output.getvalue()


def _inject_default_col_width_pt(payload: bytes) -> bytes:
    """Add the non-standard formatting attribute observed in production files."""
    output = io.BytesIO()
    with ZipFile(io.BytesIO(payload)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                marker = b"<sheetFormatPr "
                assert marker in content
                content = content.replace(marker, marker + b'defaultColWidthPt="64" ', 1)
            target.writestr(member, content)
    return output.getvalue()


@pytest.fixture(autouse=True)
def _restore_unpatched_constructor(monkeypatch):
    """Restore the pre-test constructor after each process-local shim test."""
    original_init = SheetFormatProperties.__init__
    while getattr(original_init, _MARKER, False):
        original_init = original_init.__wrapped__
    monkeypatch.setattr(SheetFormatProperties, "__init__", original_init)


def test_xlsx_conversion_shim_lets_default_col_width_pt_workbook_load():
    file_conversion._install_openpyxl_default_col_width_pt_compat()

    wrapped = getattr(SheetFormatProperties.__init__, _MARKER, False)
    assert wrapped, "openpyxl shim was not installed"

    workbook = load_workbook(io.BytesIO(_inject_default_col_width_pt(_workbook_bytes())))
    assert workbook["生产数据"]["A1"].value == "井号"
    assert workbook["生产数据"]["B1"].value == "NB27-1-1"


def test_xlsx_conversion_shim_is_idempotent():
    file_conversion._install_openpyxl_default_col_width_pt_compat()
    first_wrapper = SheetFormatProperties.__init__
    file_conversion._install_openpyxl_default_col_width_pt_compat()

    assert SheetFormatProperties.__init__ is first_wrapper


def test_xlsx_conversion_shim_preserves_standard_workbook():
    file_conversion._install_openpyxl_default_col_width_pt_compat()

    workbook = load_workbook(io.BytesIO(_workbook_bytes()))
    assert workbook["生产数据"]["B1"].value == "NB27-1-1"


def test_unpatched_openpyxl_still_reproduces_production_failure():
    # Guard: prove the fixture payload genuinely triggers the production error
    # when the shim is NOT installed, so passing tests mean the shim matters.
    with pytest.raises(TypeError, match="unexpected keyword argument 'defaultColWidthPt'"):
        load_workbook(io.BytesIO(_inject_default_col_width_pt(_workbook_bytes())))


def test_xlsx_conversion_shim_rejects_unrelated_attributes():
    file_conversion._install_openpyxl_default_col_width_pt_compat()
    with pytest.raises(TypeError, match="unrelated_attribute"):
        SheetFormatProperties(unrelated_attribute="invalid")


@pytest.mark.parametrize("with_backend_path", [False, True], ids=["without-startup-hook", "gateway-startup"])
@pytest.mark.parametrize("affected", [False, True], ids=["standard", "defaultColWidthPt"])
def test_gateway_converts_xlsx_without_startup_patch(tmp_path, affected, with_backend_path):
    source = tmp_path / "source.xlsx"
    payload = _workbook_bytes()
    if affected:
        payload = _inject_default_col_width_pt(payload)
    source.write_bytes(payload)
    env = os.environ.copy()
    env["PYTHONPATH"] = str(REPO_ROOT if with_backend_path else REPO_ROOT / "packages" / "harness")
    script = """
import asyncio
import pathlib
import sys
# Startup must not eagerly load openpyxl.
assert "openpyxl" not in sys.modules
from deerflow.utils import file_conversion
from openpyxl.worksheet.dimensions import SheetFormatProperties
assert not getattr(SheetFormatProperties.__init__, "__deerflow_default_col_width_pt_compat__", False)
# XLSX does not use this PDF option; avoid reading operator config in the test.
file_conversion._get_pdf_converter = lambda: "auto"
source = pathlib.Path(sys.argv[1])
converted = asyncio.run(file_conversion.convert_file_to_markdown(source))
assert converted is not None
assert getattr(SheetFormatProperties.__init__, "__deerflow_default_col_width_pt_compat__", False)
text = converted.read_text(encoding="utf-8")
assert "井号" in text and "NB27-1-1" in text
"""
    result = subprocess.run(
        [sys.executable, "-c", script, str(source)],
        cwd=tmp_path,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert source.read_bytes() == payload


@pytest.mark.parametrize("suffix", [".pdf", ".docx", ".pptx", ".xls", ".txt"])
def test_other_formats_do_not_install_xlsx_shim(tmp_path, monkeypatch, suffix):
    from types import SimpleNamespace
    from unittest.mock import Mock

    install = Mock(side_effect=AssertionError("XLSX shim called for another format"))
    monkeypatch.setattr(file_conversion, "_install_openpyxl_default_col_width_pt_compat", install)
    import markitdown

    monkeypatch.setattr(markitdown, "MarkItDown", lambda: SimpleNamespace(convert=lambda _: SimpleNamespace(text_content="content")))
    assert file_conversion._convert_with_markitdown(tmp_path / ("source" + suffix)) == "content"
    install.assert_not_called()


def test_xlsx_dependency_failure_uses_normal_conversion_error(tmp_path, monkeypatch, caplog):
    source = tmp_path / "source.xlsx"
    source.write_bytes(_workbook_bytes())
    monkeypatch.setattr(file_conversion, "_get_pdf_converter", lambda: "auto")

    def unavailable():
        raise ModuleNotFoundError("openpyxl unavailable", name="openpyxl")

    monkeypatch.setattr(file_conversion, "_install_openpyxl_default_col_width_pt_compat", unavailable)
    import asyncio

    assert asyncio.run(file_conversion.convert_file_to_markdown(source)) is None
    assert not source.with_suffix(".md").exists()
    assert "Failed to convert" in caplog.text
