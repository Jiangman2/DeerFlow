"""Regression tests for DeerFlow's non-invasive WeKnora DocReader XLSX shim."""

from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
COMPAT_MODULE = REPO_ROOT / "k8s" / "weknora" / "docreader-compat" / "xlsx_compat.py"


def _load_compat_module():
    spec = importlib.util.spec_from_file_location("deerflow_weknora_xlsx_compat", COMPAT_MODULE)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _workbook_bytes() -> bytes:
    output = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "生产数据"
    worksheet["A1"] = "井号"
    worksheet["B1"] = "NB27-1-1"
    workbook.save(output)
    return output.getvalue()


def _inject_default_col_width_pt(payload: bytes) -> bytes:
    """Add the non-standard formatting attribute observed in production files."""

    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                marker = b"<sheetFormatPr "
                assert marker in content
                content = content.replace(marker, marker + b'defaultColWidthPt="64" ', 1)
            target.writestr(member, content)
    return output.getvalue()


def test_unpatched_openpyxl_reproduces_production_failure():
    payload = _inject_default_col_width_pt(_workbook_bytes())

    with pytest.raises(TypeError, match="unexpected keyword argument 'defaultColWidthPt'"):
        load_workbook(BytesIO(payload))


def test_compatibility_shim_ignores_only_unsupported_formatting_attribute(monkeypatch):
    from openpyxl.worksheet.dimensions import SheetFormatProperties

    original_init = SheetFormatProperties.__init__
    monkeypatch.setattr(SheetFormatProperties, "__init__", original_init)
    compat = _load_compat_module()

    assert compat.install_openpyxl_sheet_format_compat() is True
    assert compat.install_openpyxl_sheet_format_compat() is False

    workbook = load_workbook(BytesIO(_inject_default_col_width_pt(_workbook_bytes())))
    assert workbook["生产数据"]["A1"].value == "井号"
    assert workbook["生产数据"]["B1"].value == "NB27-1-1"


def test_compatibility_shim_preserves_standard_workbooks(monkeypatch):
    from openpyxl.worksheet.dimensions import SheetFormatProperties

    original_init = SheetFormatProperties.__init__
    monkeypatch.setattr(SheetFormatProperties, "__init__", original_init)
    compat = _load_compat_module()
    compat.install_openpyxl_sheet_format_compat()

    workbook = load_workbook(BytesIO(_workbook_bytes()))
    assert workbook["生产数据"]["B1"].value == "NB27-1-1"
