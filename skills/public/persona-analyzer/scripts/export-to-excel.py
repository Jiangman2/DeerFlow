#!/usr/bin/env python3
"""
Persona-Analyzer 分析记录导出到Excel的脚本

用法:
    python export-to-excel.py <input_dir> [-o <output.xlsx>]

说明:
    - input_dir: 包含JSON分析记录的目录（或单个JSON文件）
    - 输出一个Excel文件，包含四个Sheet：
      1. 分析记录汇总 — 每条分析一行
      2. 逐节点痛点详情 — 每个痛点一行
      3. 系统性痛点分析 — 跨节点系统性模式
      4. 统计看板 — 自动生成的统计图表数据
"""

import json
import os
import sys
import argparse
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 当前沙箱镜像未预装 openpyxl；请联系平台管理员将其加入基础镜像后重试。")
    sys.exit(1)


_FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value):
    """Keep user-supplied text from becoming an Excel formula on open."""
    if not isinstance(value, str):
        return value
    return f"'{value}" if value.lstrip().startswith(_FORMULA_PREFIXES) else value


def find_json_files(path):
    """递归查找目录下所有JSON文件，或返回单个文件"""
    if os.path.isfile(path):
        if path.endswith(".json"):
            return [path]
        return []

    results = []
    for root, dirs, files in os.walk(path):
        for f in files:
            if f.endswith(".json") and f != "summary.json":
                results.append(os.path.join(root, f))
    return sorted(results)


def load_json(filepath):
    """加载JSON文件"""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, UnicodeDecodeError) as e:
        print(f"  警告: 跳过 {filepath} - JSON解析错误: {e}")
        return None


def extract_record(data, filepath):
    """从JSON记录中提取展平的数据，用于写入Excel"""
    pos = data.get("position_info", {})
    sess = data.get("session_info", {})
    bp = data.get("pain_points", {})
    stats = bp.get("summary_statistics", {})
    summary = data.get("summary", {})
    progress = data.get("progress", {})
    biz = data.get("business_flow", {})
    sys_analysis = bp.get("systemic_analysis", {})
    pri_matrix = bp.get("priority_matrix", {})

    by_type = stats.get("by_type", {})
    by_impact = stats.get("by_impact", {})

    return {
        "分析编号": data.get("analysis_id", ""),
        "分析日期": sess.get("analysis_date", ""),
        "会话状态": sess.get("session_status", ""),
        "岗位名称": pos.get("job_title", ""),
        "业务归属": pos.get("domain", ""),
        "工作地点": pos.get("location", ""),
        "常用工具": ", ".join(pos.get("tools_used", [])),
        "标签DOM": pos.get("tags", {}).get("DOM", ""),
        "标签ROL": pos.get("tags", {}).get("ROL", ""),
        "标签MOD": pos.get("tags", {}).get("MOD", ""),
        "标签DEP": pos.get("tags", {}).get("DEP", ""),
        "标签LVL": pos.get("tags", {}).get("LVL", ""),
        "标签DAT": pos.get("tags", {}).get("DAT", ""),
        "标签SKL": ", ".join(pos.get("tags", {}).get("SKL", [])),
        "一句话总结": pos.get("summary", ""),
        "业务描述": biz.get("work_description", ""),
        "流程节点数": len(biz.get("nodes", [])),
        "自定义节点数": sum(1 for n in biz.get("nodes", []) if n.get("is_custom", False)),
        "已生成SVG流程图": "是" if biz.get("flowchart_svg_generated", False) else "否",
        "已分析节点数": stats.get("total_nodes_analyzed", 0),
        "剩余节点数": stats.get("total_nodes_remaining", 0),
        "痛点总数": stats.get("total_pain_points", 0),
        "技术痛点": by_type.get("技术痛点", 0),
        "流程痛点": by_type.get("流程痛点", 0),
        "数据痛点": by_type.get("数据痛点", 0),
        "协同痛点": by_type.get("协同痛点", 0),
        "管理痛点": by_type.get("管理痛点", 0),
        "高影响": by_impact.get("高", 0),
        "中影响": by_impact.get("中", 0),
        "低影响": by_impact.get("低", 0),
        "Quick Wins数": len(pri_matrix.get("quick_wins", [])),
        "重点攻克数": len(pri_matrix.get("focus_areas", [])),
        "长期规划数": len(pri_matrix.get("long_term", [])),
        "存在系统性痛点": "是" if sys_analysis.get("has_systemic_issues", False) else "否",
        "系统性痛点模式数": len(sys_analysis.get("patterns", [])),
        "对话轮次": sess.get("total_turns", 0),
        "时长(分钟)": sess.get("total_duration_minutes", 0),
        "疲劳事件数": len(sess.get("fatigue_events", [])),
        "情绪触发疲劳数": sum(1 for e in sess.get("fatigue_events", []) if e.get("trigger") == "emotion_word"),
        "已完成阶段": ", ".join(progress.get("completed_stages", [])),
        "已完成节点": ", ".join(progress.get("completed_nodes", [])),
        "剩余待分析节点": ", ".join(progress.get("remaining_nodes", [])),
        "下次续接": progress.get("next_resume_point", ""),
        "关键发现": "\n".join(summary.get("key_findings", [])),
        "系统性建议": "\n".join(summary.get("systemic_recommendations", [])),
        "原始文件": os.path.basename(filepath),
    }


def extract_pain_point_details(data):
    """提取逐节点痛点详情"""
    rows = []
    pos = data.get("position_info", {})
    job_title = pos.get("job_title", "")
    domain = pos.get("domain", "")
    analysis_id = data.get("analysis_id", "")

    bp = data.get("pain_points", {})
    details = bp.get("details", [])

    for node in details:
        node_name = node.get("node_name", "")
        node_id = node.get("node_id", "")
        for pp in node.get("pain_points_list", []):
            root_cause_chain = pp.get("root_cause_chain", [])
            rows.append({
                "分析编号": analysis_id,
                "岗位名称": job_title,
                "业务归属": domain,
                "节点序号": node_id,
                "节点名称": node_name,
                "痛点描述": pp.get("description", ""),
                "痛点类型": pp.get("type", ""),
                "影响程度": pp.get("impact", ""),
                "解决难度": pp.get("difficulty", ""),
                "发生频率": pp.get("frequency", ""),
                "当前做法": pp.get("current_approach", ""),
                "理想状态": pp.get("ideal_state", ""),
                "已应用5 Why": "是" if pp.get("five_why_applied", False) else "否",
                "根因链": " → ".join(root_cause_chain) if root_cause_chain else "",
                "最终根因": pp.get("root_cause", ""),
                "对话摘录": pp.get("dialogue_excerpt", ""),
            })

    return rows


def extract_systemic_analysis(data):
    """提取系统性痛点关联分析"""
    rows = []
    pos = data.get("position_info", {})
    job_title = pos.get("job_title", "")
    domain = pos.get("domain", "")
    analysis_id = data.get("analysis_id", "")

    bp = data.get("pain_points", {})
    sys_analysis = bp.get("systemic_analysis", {})

    if not sys_analysis.get("has_systemic_issues", False):
        return rows

    for pattern in sys_analysis.get("patterns", []):
        rows.append({
            "分析编号": analysis_id,
            "岗位名称": job_title,
            "业务归属": domain,
            "痛点类型": pattern.get("pain_type", ""),
            "受影响节点": ", ".join(pattern.get("affected_nodes", [])),
            "受影响节点数": len(pattern.get("affected_nodes", [])),
            "模式描述": pattern.get("pattern_description", ""),
            "系统性根因假设": pattern.get("systemic_root_cause_hypothesis", ""),
        })

    return rows


def get_summary_statistics(records):
    """生成汇总统计"""
    from collections import Counter

    total_analyses = len(records)

    # 岗位分布
    job_counter = Counter(r.get("岗位名称", "") for r in records)
    # 业务域分布
    domain_counter = Counter(r.get("业务归属", "") for r in records)
    # 状态分布
    status_counter = Counter(r.get("会话状态", "") for r in records)
    # 痛点类型统计
    type_sum = {
        "技术痛点": sum(r.get("技术痛点", 0) for r in records),
        "流程痛点": sum(r.get("流程痛点", 0) for r in records),
        "数据痛点": sum(r.get("数据痛点", 0) for r in records),
        "协同痛点": sum(r.get("协同痛点", 0) for r in records),
        "管理痛点": sum(r.get("管理痛点", 0) for r in records),
    }
    # 优先级矩阵统计
    qw_total = sum(r.get("Quick Wins数", 0) for r in records)
    focus_total = sum(r.get("重点攻克数", 0) for r in records)
    long_total = sum(r.get("长期规划数", 0) for r in records)
    # 系统性痛点统计
    sys_count = sum(1 for r in records if r.get("存在系统性痛点") == "是")
    # 5 Why 统计
    five_why_total = 0  # 需要在 pain point details 中统计

    stats = []
    stats.append({"统计项": "总分析次数", "数值": total_analyses})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【会话状态分布】", "数值": ""})
    for s, c in status_counter.most_common():
        stats.append({"统计项": f"  {s}", "数值": c})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【岗位分布】", "数值": ""})
    for j, c in job_counter.most_common():
        stats.append({"统计项": f"  {j}", "数值": c})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【业务域分布】", "数值": ""})
    for d, c in domain_counter.most_common():
        stats.append({"统计项": f"  {d}", "数值": c})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【痛点类型汇总】", "数值": ""})
    for t, c in type_sum.items():
        stats.append({"统计项": f"  {t}", "数值": c})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【优先级矩阵分布】", "数值": ""})
    stats.append({"统计项": "  Quick Wins（高影响+易解决）", "数值": qw_total})
    stats.append({"统计项": "  重点攻克（高影响+中难度）", "数值": focus_total})
    stats.append({"统计项": "  长期规划（高影响+难解决）", "数值": long_total})
    stats.append({"统计项": "---", "数值": ""})
    stats.append({"统计项": "【系统性分析】", "数值": ""})
    stats.append({"统计项": "  存在系统性痛点的分析数", "数值": sys_count})

    return stats


def style_header(ws, row, max_col, fill_color="1F4E79", font_color="FFFFFF"):
    """给表头行添加样式"""
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_font = Font(name="微软雅黑", bold=True, color=font_color, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border


def auto_width(ws, min_width=8, max_width=40):
    """自动调整列宽"""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                # 估算中文字符宽度
                val = str(cell.value)
                char_len = 0
                for ch in val:
                    if ord(ch) > 127:
                        char_len += 2
                    else:
                        char_len += 1
                max_len = max(max_len, char_len)
        adjusted = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


def export_to_excel(input_path, output_path=None):
    """导出分析记录到Excel"""
    if output_path is None:
        output_path = f"persona-analyzer-export-{datetime.now().strftime('%Y%m%d')}.xlsx"

    print(f"📂 扫描路径: {input_path}")
    files = find_json_files(input_path)
    print(f"📄 找到 {len(files)} 个JSON分析记录文件")

    if not files:
        print("❌ 未找到任何JSON文件，请检查路径")
        return

    # 加载所有记录
    records = []
    all_pain_rows = []
    all_sys_rows = []
    for fp in files:
        data = load_json(fp)
        if data:
            records.append(extract_record(data, fp))
            all_pain_rows.extend(extract_pain_point_details(data))
            all_sys_rows.extend(extract_systemic_analysis(data))

    print(f"✅ 成功加载 {len(records)} 条分析记录")
    print(f"📊 提取 {len(all_pain_rows)} 条痛点详情")
    print(f"🔗 提取 {len(all_sys_rows)} 条系统性关联分析")

    # 创建Excel工作簿
    wb = openpyxl.Workbook()

    # ── Sheet 1: 分析记录汇总 ──
    ws1 = wb.active
    ws1.title = "分析记录汇总"

    if records:
        headers = list(records[0].keys())
        # 写表头
        for col, h in enumerate(headers, 1):
            ws1.cell(row=1, column=col, value=h)
        style_header(ws1, 1, len(headers))

        # 写数据
        data_font = Font(name="微软雅黑", size=10)
        data_align = Alignment(vertical="center", wrap_text=True)
        for row_idx, rec in enumerate(records, 2):
            for col_idx, h in enumerate(headers, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=safe_excel_value(rec.get(h, "")))
                cell.font = data_font
                cell.alignment = data_align

        auto_width(ws1)
        # 冻结首行
        ws1.freeze_panes = "A2"

    # ── Sheet 2: 逐节点痛点详情 ──
    ws2 = wb.create_sheet("逐节点痛点详情")

    if all_pain_rows:
        headers2 = list(all_pain_rows[0].keys())
        for col, h in enumerate(headers2, 1):
            ws2.cell(row=1, column=col, value=h)
        style_header(ws2, 1, len(headers2))

        data_font = Font(name="微软雅黑", size=10)
        data_align = Alignment(vertical="center", wrap_text=True)
        for row_idx, pr in enumerate(all_pain_rows, 2):
            for col_idx, h in enumerate(headers2, 1):
                cell = ws2.cell(row=row_idx, column=col_idx, value=safe_excel_value(pr.get(h, "")))
                cell.font = data_font
                cell.alignment = data_align

        auto_width(ws2, max_width=60)
        ws2.freeze_panes = "A2"

    # ── Sheet 3: 系统性痛点分析 ──
    ws3 = wb.create_sheet("系统性痛点分析")
    sheet3_title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    sheet3_title_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)

    if all_sys_rows:
        headers3 = list(all_sys_rows[0].keys())
        for col, h in enumerate(headers3, 1):
            ws3.cell(row=1, column=col, value=h)
        style_header(ws3, 1, len(headers3))

        data_font = Font(name="微软雅黑", size=10)
        data_align = Alignment(vertical="center", wrap_text=True)
        for row_idx, sr in enumerate(all_sys_rows, 2):
            for col_idx, h in enumerate(headers3, 1):
                cell = ws3.cell(row=row_idx, column=col_idx, value=safe_excel_value(sr.get(h, "")))
                cell.font = data_font
                cell.alignment = data_align

        auto_width(ws3, max_width=60)
        ws3.freeze_panes = "A2"
    else:
        ws3.cell(row=1, column=1, value="暂无跨节点的系统性痛点数据")
        ws3.cell(row=1, column=1).font = Font(name="微软雅黑", size=12, color="666666")
        ws3.column_dimensions["A"].width = 40

    # ── Sheet 4: 统计看板 ──
    ws4 = wb.create_sheet("统计看板")

    stats_data = get_summary_statistics(records)
    headers4 = ["统计项", "数值"]
    for col, h in enumerate(headers4, 1):
        ws4.cell(row=1, column=col, value=h)
    style_header(ws4, 1, 2)

    data_font = Font(name="微软雅黑", size=10)
    section_font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
    for row_idx, s in enumerate(stats_data, 2):
        for col_idx, h in enumerate(headers4, 1):
            cell = ws4.cell(row=row_idx, column=col_idx, value=safe_excel_value(s.get(h, "")))
            cell.font = section_font if s["统计项"].startswith("【") else data_font
            cell.alignment = Alignment(vertical="center")

    ws4.column_dimensions["A"].width = 35
    ws4.column_dimensions["B"].width = 15
    ws4.freeze_panes = "A2"

    # 保存
    wb.save(output_path)
    print(f"\n✅ 导出成功!")
    print(f"📁 输出文件: {os.path.abspath(output_path)}")
    print(f"   📊 Sheet 1: 分析记录汇总 ({len(records)}条)")
    print(f"   📋 Sheet 2: 逐节点痛点详情 ({len(all_pain_rows)}条)")
    print(f"   🔗 Sheet 3: 系统性痛点分析 ({len(all_sys_rows)}条)")
    print(f"   📈 Sheet 4: 统计看板")


def main():
    parser = argparse.ArgumentParser(
        description="Persona-Analyzer 分析记录导出到Excel",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python export-to-excel.py /data/analysis-records/
  python export-to-excel.py ./analysis.json -o my-report.xlsx
        """
    )
    parser.add_argument("input", help="JSON分析记录目录或单个JSON文件路径")
    parser.add_argument("-o", "--output", help="输出Excel文件路径（可选）")

    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"错误: 路径不存在: {args.input}")
        sys.exit(1)

    export_to_excel(args.input, args.output)


if __name__ == "__main__":
    main()
